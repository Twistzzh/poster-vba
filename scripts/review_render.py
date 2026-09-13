# -*- coding: utf-8 -*-
"""poster-vba 第二轮评审的程序化几何检查（对 build_poster.py 产出的 .pptx）。

用法:
    python review_render.py <输出目录|poster.pptx路径>

检查项（对应 references/review_rubric.md 评审二）:
  1. 文本溢出    PIL 按真实字体量宽（无字体文件时退化为系数估算），
                 文字溢出所属面板/节点框即 FAIL；单行超宽但放得下高度 ->
                 依赖自动折行，WARN。
  2. 字号观感    所有节点"文本宽/可用宽"中位数过低 -> 提示字体可能过小。
  3. 图标/图片  出现灰色虚线 "IMG 文件名" 占位框 -> FAIL（素材缺失）；
                 图片超出幻灯片边界 -> FAIL；图片过小可能发糊 -> WARN。
  4. 装饰遮挡    无文字的卡通/装饰形状大面积覆盖有文字节点 -> WARN。
  5. 出界        形状超出幻灯片边界 -> FAIL。

退出码: 0 = PASS（允许 WARN），1 = FAIL。文本型结论打印到 stdout。
"""
import math
import os
import sys

FALLBACK_K = 0.55   # 无字体文件时的估算系数（em/字符）
K_LINE = 1.25       # 行高系数
TOL_EDGE = 0.06     # 端点贴边容差（英寸）
CONTAINER_MIN = 3   # bbox 内包含 >=3 个其他节点中心 -> 视为容器面板
OVERLAP_WARN = 0.45  # 装饰遮挡有文字节点的交叠比例阈值

FONT_FILES = {
    "microsoft yahei": "msyh.ttc", "yahei": "msyh.ttc", "msyh": "msyh.ttc",
    "arial": "arial.ttf", "helvetica": "arial.ttf", "calibri": "calibri.ttf",
    "times new roman": "times.ttf", "times": "times.ttf",
    "simsun": "simsun.ttc", "宋体": "simsun.ttc", "simhei": "simhei.ttf", "黑体": "simhei.ttf",
}
_font_cache = {}
_pil_ok = None


def _load_font(font_name, pt, bold=False):
    global _pil_ok
    if _pil_ok is False:
        return None
    try:
        from PIL import ImageFont
    except ImportError:
        _pil_ok = False
        return None
    key = (font_name.lower(), int(pt * 4), bold)
    if key in _font_cache:
        return _font_cache[key]
    fn = FONT_FILES.get(font_name.lower())
    if not fn:
        _pil_ok = False
        return None
    if bold and fn == "msyh.ttc":
        bpath = os.path.join(os.environ.get("WINDIR", r"C:\Windows"),
                             "Fonts", "msyhbd.ttc")
        if os.path.isfile(bpath):
            fn = "msyhbd.ttc"
    path = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts", fn)
    if not os.path.isfile(path):
        _pil_ok = False
        return None
    try:
        f = ImageFont.truetype(path, size=max(8, int(pt * 4)))
    except Exception:
        _pil_ok = False
        return None
    _font_cache[key] = f
    return f


def _cjk_subst(font_name):
    f = (font_name or "").lower()
    if any(k in f for k in ("yahei", "msyh", "arial", "helvetica", "segoe", "sans")):
        return "microsoft yahei"
    if any(k in f for k in ("simsun", "宋", "sun", "serif", "times")):
        return "simsun"
    if any(k in f for k in ("simhei", "黑", "hei")):
        return "simhei"
    return None


def is_cjk(ch):
    return ('\u3000' <= ch <= '\u303f' or '\u3400' <= ch <= '\u9fff'
            or '\uff00' <= ch <= '\uffef')


def text_width_in(text, font_name, pt, bold=False):
    f = _load_font(font_name, pt, bold)
    if f is None and any(is_cjk(c) for c in text):
        sub = _cjk_subst(font_name)
        if sub:
            f = _load_font(sub, pt, bold)
    if f is not None:
        try:
            return f.getlength(text) / 4.0 / 72.0
        except Exception:
            pass
    w = 0.0
    for ch in text:
        if ch == ' ':
            w += 0.3 * pt
        elif is_cjk(ch):
            w += 1.0 * pt
        else:
            w += FALLBACK_K * pt
    return w / 72.0


def emu2in(v):
    return v / 914400.0


def bbox(shp):
    return (emu2in(shp.left), emu2in(shp.top),
            emu2in(shp.left + shp.width), emu2in(shp.top + shp.height))


def inside(px, py, b, tol=TOL_EDGE):
    x0, y0, x1, y1 = b
    return (x0 - tol <= px <= x1 + tol) and (y0 - tol <= py <= y1 + tol)


def text_metrics(tf, default_pt=18.0):
    lines, sizes, fonts = [], [], []
    for para in tf.paragraphs:
        t = "".join(r.text for r in para.runs)
        if not t and para.runs:
            t = " "
        if t:
            lines.append(t)
        for r in para.runs:
            if r.font.size is not None:
                sizes.append(r.font.size.pt)
            if r.font.name:
                fonts.append(r.font.name)
    pt = max(sizes) if sizes else None
    font = fonts[0] if fonts else "Microsoft YaHei"
    return lines, (pt or default_pt), font, pt is not None


def review(pptx_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(pptx_path)
    errs, warns, infos = [], [], []
    slide = prs.slides[0]
    SW, SH = emu2in(prs.slide_width), emu2in(prs.slide_height)

    nodes, pics = [], []
    for shp in slide.shapes:
        is_pic = (shp.shape_type == MSO_SHAPE_TYPE.PICTURE)
        b = bbox(shp)
        tf = shp.text_frame if shp.has_text_frame else None
        label = tf.text.replace("\n", " / ")[:40] if tf is not None else ""
        rec = {"id": shp.shape_id, "name": "%s [%s]" % (shp.shape_id, label),
               "b": b, "tf": tf, "text": bool(label.strip())}
        if is_pic:
            pics.append(rec)
        else:
            nodes.append(rec)

    # 背景整版矩形（覆盖整张幻灯片）不计为面板/文字块
    def is_full_bg(n):
        x0, y0, x1, y1 = n["b"]
        return (x0 <= -0.01 and y0 <= -0.01 and x1 >= SW - 0.01 and y1 >= SH - 0.01)

    bg = [n for n in nodes if is_full_bg(n)]
    solids = [n for n in nodes if not is_full_bg(n)]

    # 缺失素材占位框（build_poster 的回放/COM 兜底）：文字以 "IMG " 开头
    for n in solids:
        if n["tf"] is not None and n["tf"].text.strip().startswith("IMG "):
            errs.append("%s: 图标素材缺失，渲染成占位框 %r（确认 ASSET_DIR 与文件名）"
                        % (n["name"], n["tf"].text.strip()))

    # ---- 1) 文本溢出 + 字号观感 -------------------------------------------
    ratios = []
    for n in solids:
        if n["tf"] is None or not n["tf"].text.strip():
            continue
        tf = n["tf"]
        lines, pt, font, known = text_metrics(tf)
        if not lines:
            continue
        if not known:
            warns.append("%s: 字号未显式设置（读不到 run.font.size）" % n["name"])
        ml = emu2in(tf.margin_left) if tf.margin_left is not None else 0.05
        mr = emu2in(tf.margin_right) if tf.margin_right is not None else 0.05
        mt = emu2in(tf.margin_top) if tf.margin_top is not None else 0.02
        mb = emu2in(tf.margin_bottom) if tf.margin_bottom is not None else 0.02
        avail_w = (n["b"][2] - n["b"][0]) - (ml + mr)
        avail_h = (n["b"][3] - n["b"][1]) - (mt + mb)
        if avail_w <= 0 or avail_h <= 0:
            continue
        line_ws = [text_width_in(l, font, pt) for l in lines]
        h_coeff = 1.0 if len(lines) == 1 else K_LINE
        total_h = len(lines) * pt * h_coeff / 72.0
        # 只统计多行或较高的正文块：短标题(单行小条)与图表数值/分类小标签
        # 天生文字窄，纳入中位数会误报"字体偏小"。真问题(正文块字号过小)仍保留。
        box_h = n["b"][3] - n["b"][1]
        if len(lines) >= 2 or box_h >= 0.6:
            ratios.append(max(line_ws) / avail_w)
        over = [(l, w) for l, w in zip(lines, line_ws) if w > avail_w]
        if over:
            if len(lines) == 1 and total_h <= avail_h:
                wrap_n = int(math.ceil(line_ws[0] / avail_w))
                if wrap_n * pt * K_LINE / 72.0 > avail_h:
                    errs.append("%s: 字号 %gpt 过大，'%s' 自动折行 %d 行仍溢出面板"
                                % (n["name"], pt, over[0][0][:30], wrap_n))
                else:
                    warns.append("%s: '%s' 超宽依赖自动折行（应用 vbLf 显式分行）"
                                 % (n["name"], over[0][0][:30]))
            else:
                for l, w in over:
                    errs.append("%s: 字号 %gpt 文字 '%s' 溢出面板宽（est %.2f in > 可用 %.2f in）"
                                % (n["name"], pt, l[:30], w, avail_w))
        if total_h > avail_h + 0.02:
            tiny = ((n["b"][3] - n["b"][1]) < 0.25 and len(lines) == 1
                    and len("".join(lines).strip()) <= 4)
            if tiny:
                warns.append("%s: 小值标签 %.2f in 略高于 %.2f in 的框，检查是否仍在视觉可接受范围"
                             % (n["name"], total_h, avail_h))
            else:
                errs.append("%s: 文本总高 %.2f in 超出面板可用高 %.2f in（字号过大或行数过多）"
                            % (n["name"], total_h, avail_h))
    if ratios:
        med = sorted(ratios)[len(ratios) // 2]
        if med < 0.35:
            warns.append("文字宽/面板宽中位数 %.2f < 0.35：字体可能整体偏小（对照原图比例）" % med)
        else:
            infos.append("文字宽/面板宽中位数 %.2f（0.35~1.0 为合理区间）" % med)

    # ---- 3) 图片：出界 / 过小 --------------------------------------------
    for p in pics:
        b = p["b"]
        if b[0] < -0.01 or b[1] < -0.01 or b[2] > SW + 0.01 or b[3] > SH + 0.01:
            errs.append("图标/图片 %s 超出幻灯片边界" % p["name"])
        w_in = b[2] - b[0]
        if w_in < 0.4:
            warns.append("图标/图片 %s 宽度仅 %.2f in，放大后可能发糊" % (p["name"], w_in))

    # ---- 4) 装饰遮挡（无文字形状大面积盖有文字节点） -----------------------
    decos = [n for n in solids if not n["text"]]
    for dnode in decos:
        bx = dnode["b"]
        area = (bx[2] - bx[0]) * (bx[3] - bx[1])
        if area <= 0:
            continue
        for n in solids:
            if not n["text"] or n is dnode:
                continue
            ox = min(bx[2], n["b"][2]) - max(bx[0], n["b"][0])
            oy = min(bx[3], n["b"][3]) - max(bx[1], n["b"][1])
            if ox <= 0 or oy <= 0:
                continue
            # 包含关系（装饰在面板内，正常）跳过
            if (bx[0] <= n["b"][0] and bx[1] <= n["b"][1]
                    and bx[2] >= n["b"][2] and bx[3] >= n["b"][3]):
                continue
            frac = (ox * oy) / area
            if frac >= OVERLAP_WARN:
                warns.append("装饰 %s 覆盖文字节点 %s 约 %.0f%%（挪开，别遮文字）"
                             % (dnode["name"], n["name"], 100 * frac))

    # ---- 5) 出界 ------------------------------------------------------------
    for n in nodes + pics:
        b = n["b"]
        if b[0] < -0.01 or b[1] < -0.01 or b[2] > SW + 0.01 or b[3] > SH + 0.01:
            if n in bg:
                continue
            errs.append("形状 %s 超出幻灯片边界 %s" % (n["name"], (round(SW, 2), round(SH, 2))))
    return errs, warns, infos, len(nodes), len(pics)


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    target = sys.argv[1]
    if os.path.isdir(target):
        cand = [f for f in os.listdir(target)
                if f.endswith(".pptx") and not f.startswith("~$")]
        if not cand:
            print("FAIL: 目录里没有 .pptx，请先跑 build_poster.py")
            return 1
        target = os.path.join(target, max(cand, key=lambda f: os.path.getmtime(os.path.join(target, f))))
    if not os.path.isfile(target):
        print("FAIL: 找不到 %s" % target)
        return 1
    print("review_render: %s" % target)
    try:
        errs, warns, infos, nn, npic = review(target)
    except ImportError:
        print("FAIL: 需要 python-pptx，请用隔离 venv 运行:")
        print("  C:/Users/Hello/.workbuddy/binaries/python/envs/default/Scripts/python.exe "
              + sys.argv[0])
        return 2
    print("shapes: %d 文本/面板块 / %d 图片" % (nn, npic))
    for line in errs:
        print("  [ERR ] %s" % line)
    for line in warns:
        print("  [WARN] %s" % line)
    for line in infos:
        print("  [INFO] %s" % line)
    if errs:
        print("RESULT: FAIL (%d err / %d warn)" % (len(errs), len(warns)))
        print("按 references/review_rubric.md 评审二处置：修 .bas -> 重跑 build_poster.py -> 再评审")
        return 1
    print("RESULT: PASS (%d warn)" % len(warns))
    print("程序化检查通过。仍需多模态对照 preview_compare.png 做风格终审（评审一 + 评审二目测项）。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
